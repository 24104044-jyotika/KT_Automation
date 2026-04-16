import os, logging, json, secrets
from functools import wraps
from datetime import datetime, timedelta
 
from flask import (Flask, render_template, redirect, url_for,
                   flash, request, jsonify, send_from_directory, make_response, session)
from flask_login import (LoginManager, login_user, logout_user,
                         login_required, current_user)
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
 
from config import Config
 
logging.basicConfig(level=logging.INFO)
app = Flask(__name__)
app.config.from_object(Config)
 
# ── IMPORTANT: init extensions BEFORE importing models ──
from extensions import db, mail
db.init_app(app)
# mail.init_app(app)  # Removed: Brevo HTTP API used instead of Flask-Mail
 
# ── NOW import models (db is ready) ──
from models import *
 
from flask_sqlalchemy import SQLAlchemy
# from flask_mail import Mail  # Removed: using Brevo HTTP API
 
os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)
os.makedirs(os.path.join(app.config['UPLOAD_FOLDER'], 'marksheets'), exist_ok=True)
os.makedirs(os.path.join(app.config['UPLOAD_FOLDER'], 'caste_certs'), exist_ok=True)
 
login_manager = LoginManager(app)
login_manager.login_view = 'login'

# ─── Fee Constants ────────────────────────────────────────────────────────────
REVAL_FEE_PER_PAPER  = 200   # ₹200 per paper
PHOTOCOPY_FEE        = 200   # ₹200 flat per paper

# ─── Semester ↔ Subject Mapping (MU Engineering – common subjects) ─────────────
SEMESTER_SUBJECTS = {
    '1': [
        'Applied Mathematics I',
        'Engineering Physics I',
        'Engineering Chemistry I',
        'Engineering Mechanics',
        'Basic Electrical Engineering',
        'Engineering Graphics',
        'Professional Communication & Ethics I',
    ],
    '2': [
        'Applied Mathematics II',
        'Engineering Physics II',
        'Engineering Chemistry II',
        'Engineering Drawing',
        'Basic Electronics Engineering',
        'Environmental Studies',
        'Professional Communication & Ethics II',
    ],
    '3': [
        'Applied Mathematics III',
        'Data Structures',
        'Digital Logic Design',
        'Discrete Mathematics',
        'Computer Organization & Architecture',
        'Object Oriented Programming',
        'Analog & Digital Communication',
    ],
    '4': [
        'Engineering Mathematics IV',
        'Analysis of Algorithms',
        'Database Management Systems',
        'Operating Systems',
        'Computer Networks',
        'Microprocessor',
        'Software Engineering',
    ],
    '5': [
        'Internet Programming',
        'Computer Graphics',
        'Theoretical Computer Science',
        'System Programming & Compiler Construction',
        'Information & Network Security',
        'Mobile Communication',
    ],
    '6': [
        'Distributed Computing',
        'Machine Learning',
        'Cloud Computing',
        'Big Data Analytics',
        'Information Retrieval',
        'Digital Signal Processing',
        'User Interface Design',
    ],
    '7': [
        'Artificial Intelligence',
        'Deep Learning',
        'Blockchain Technology',
        'Internet of Things',
        'Data Warehousing & Mining',
        'Software Project Management',
        'Technical Seminar',
    ],
    '8': [
        'Project Management & Entrepreneurship',
        'Audit Course',
        'Major Project',
        'Elective I',
        'Elective II',
    ],
}




# ─── Helpers ──────────────────────────────────────────────────────────────────

@login_manager.user_loader
def load_user(uid): return User.query.get(int(uid))

def add_notification(user_id, title, message, notif_type='info'):
    n = Notification(user_id=user_id, title=title, message=message, notif_type=notif_type)
    db.session.add(n); db.session.commit()

def allowed_file(filename):
    return ('.' in filename and
            filename.rsplit('.', 1)[1].lower() in app.config['ALLOWED_EXTENSIONS'])

def notify_student(student, title, message, notif_type,
                   email_fn=None, email_args=(),
                   wa_fn=None,    wa_args=()):
    add_notification(student.id, title, message, notif_type)
    if email_fn:
        try: email_fn(mail, *email_args)
        except Exception as e: app.logger.error(f"[Email] {e}")
    if wa_fn and student.phone:
        try: wa_fn(app, student.phone, *wa_args)
        except Exception as e: app.logger.error(f"[WhatsApp] {e}")

def admin_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if not current_user.is_authenticated or current_user.role != 'admin':
            flash('Admin access required.', 'error')
            return redirect(url_for('login'))
        return f(*args, **kwargs)
    return decorated

# ─── Redirects ────────────────────────────────────────────────────────────────

@app.route('/admin')
def admin_redirect(): 
    return redirect(url_for('admin_dashboard'))

@app.route('/student')
def student_redirect(): 
    return redirect(url_for('student_dashboard'))

# ─── Auth ─────────────────────────────────────────────────────────────────────

@app.route('/')
def index():
    if current_user.is_authenticated:
        return redirect(url_for('admin_dashboard' if current_user.role == 'admin'
                                else 'student_dashboard'))
    return redirect(url_for('login'))

@app.route('/login', methods=['GET','POST'])
def login():
    if current_user.is_authenticated: 
        return redirect(url_for('index'))
    if request.method == 'POST':
        user = User.query.filter_by(email=request.form.get('email','').strip()).first()
        if user and user.check_password(request.form.get('password','')):
            # Skip OTP for admin users, direct login
            if user.role == 'admin':
                login_user(user)
                return redirect(url_for('index'))
            
            # Generate OTP and send via email for students
            import random
            otp = ''.join([str(random.randint(0, 9)) for _ in range(6)])
            user.set_otp(otp)
            db.session.commit()
            try:
                from email_utils import send_login_otp
                send_login_otp(mail, user.name, user.email, otp)
            except Exception as e:
                app.logger.error(f"[OTP Email] {e}")
                flash('Error sending OTP. Please try again.', 'error')
                return render_template('login.html')
            # Store email in session for OTP verification
            session['otp_email'] = user.email
            flash('OTP has been sent to your email. Please verify to complete login.', 'info')
            return redirect(url_for('verify_otp'))
        flash('Invalid email or password.', 'error')
    return render_template('login.html')

@app.route('/register', methods=['GET','POST'])
def register():
    if current_user.is_authenticated: 
        return redirect(url_for('index'))
    if request.method == 'POST':
        name       = request.form.get('name','').strip()
        email      = request.form.get('email','').strip()
        phone      = request.form.get('phone','').strip()
        department = request.form.get('department','').strip()
        semester   = request.form.get('semester','').strip()
        pw         = request.form.get('password','')
        if User.query.filter_by(email=email).first():
            flash('Email already registered.', 'error')
            return render_template('register.html')

        # ── Optional caste certificate ──────────────────────────────────────
        caste_cert_path = ''
        cc_file = request.files.get('caste_certificate')
        if cc_file and cc_file.filename:
            if allowed_file(cc_file.filename):
                cc_name = f"cc_{email.replace('@','_')}_{datetime.utcnow().strftime('%Y%m%d%H%M%S')}_{secure_filename(cc_file.filename)}"
                cc_save  = os.path.join(app.config['UPLOAD_FOLDER'], 'caste_certs', cc_name)
                cc_file.save(cc_save)
                caste_cert_path = cc_name
            else:
                flash('Caste certificate must be PNG, JPG, or PDF.', 'error')
                return render_template('register.html')

        user = User(name=name, email=email, phone=phone, department=department,
                    semester=semester, role='student',
                    caste_certificate=caste_cert_path)
        user.set_password(pw)
        db.session.add(user)
        db.session.commit()

        # ── Link this email to any existing FailedStudentRecord rows ────────
        from models import FailedStudentRecord
        unlinked = FailedStudentRecord.query.filter(
            FailedStudentRecord.email == '',
            FailedStudentRecord.user_id == None
        ).all()
        # also check records that matched on name but had no email
        name_matched = FailedStudentRecord.query.filter(
            FailedStudentRecord.student_name.ilike(f'%{name}%'),
            FailedStudentRecord.email == ''
        ).all()
        for rec in set(unlinked) | set(name_matched):
            if rec.student_name.lower() in name.lower() or name.lower() in rec.student_name.lower():
                rec.email   = email
                rec.user_id = user.id
        # also update any records already matched by seat/prn but missing email
        FailedStudentRecord.query.filter_by(user_id=user.id, email='').update({'email': email})
        db.session.commit()

        from email_utils import send_welcome
        from whatsapp_utils import wa_welcome
        notify_student(user, 'Welcome!', f'Hello {name}, your account is ready.', 'success',
            email_fn=send_welcome, email_args=(name, email),
            wa_fn=wa_welcome,      wa_args=(name,))
        flash('Registration successful! Please log in with your credentials.', 'success')
        return redirect(url_for('login'))
    return render_template('register.html')

@app.route('/verify-otp', methods=['GET','POST'])
def verify_otp():
    if current_user.is_authenticated: return redirect(url_for('index'))
    if 'otp_email' not in session:
        flash('Please log in first.', 'error')
        return redirect(url_for('login'))
    
    if request.method == 'POST':
        otp_code = request.form.get('otp','').strip()
        user = User.query.filter_by(email=session['otp_email']).first()
        
        if not user:
            flash('User not found.', 'error')
            return redirect(url_for('login'))
        
        if user.verify_otp(otp_code):
            # OTP is valid, clear it and log in the user
            user.clear_otp()
            db.session.commit()
            login_user(user)
            session.pop('otp_email', None)
            flash('Login successful!', 'success')
            return redirect(url_for('index'))
        else:
            flash('Invalid or expired OTP. Please try again.', 'error')
    
    return render_template('verify_otp.html', email=session.get('otp_email', ''))

@app.route('/logout')
@login_required
def logout():
    logout_user(); return redirect(url_for('login'))

# ─── Password Reset ───────────────────────────────────────────────────────────

@app.route('/forgot-password', methods=['GET','POST'])
def forgot_password():
    if current_user.is_authenticated: return redirect(url_for('index'))
    if request.method == 'POST':
        email = request.form.get('email','').strip()
        user = User.query.filter_by(email=email).first()
        if user:
            token = secrets.token_urlsafe(32)
            user.reset_token = token
            user.reset_token_expiry = datetime.utcnow() + timedelta(hours=1)
            db.session.commit()
            try:
                from email_utils import send_password_reset
                send_password_reset(mail, user.name, email, token)
            except Exception as e:
                app.logger.error(f"[Password Reset Email] {e}")
        # Always show success to prevent user enumeration
        flash('If that email exists, a password reset link has been sent. Check your inbox.', 'success')
        return redirect(url_for('login'))
    return render_template('forgot_password.html')

@app.route('/reset-password/<token>', methods=['GET','POST'])
def reset_password(token):
    if current_user.is_authenticated: return redirect(url_for('index'))
    user = User.query.filter_by(reset_token=token).first()
    if not user or not user.reset_token_expiry or user.reset_token_expiry < datetime.utcnow():
        flash('This password reset link is invalid or has expired.', 'error')
        return redirect(url_for('forgot_password'))
    if request.method == 'POST':
        pw  = request.form.get('password','')
        pw2 = request.form.get('confirm_password','')
        if pw != pw2:
            flash('Passwords do not match.', 'error')
            return render_template('reset_password.html', token=token)
        if len(pw) < 6:
            flash('Password must be at least 6 characters.', 'error')
            return render_template('reset_password.html', token=token)
        user.set_password(pw)
        user.reset_token = None
        user.reset_token_expiry = None
        db.session.commit()
        flash('Password reset successfully! Please log in.', 'success')
        return redirect(url_for('login'))
    return render_template('reset_password.html', token=token)

# ─── Student Dashboard ────────────────────────────────────────────────────────

@app.route('/student/dashboard')
@login_required
def student_dashboard():
    if current_user.role == 'admin': return redirect(url_for('admin_dashboard'))
    active_rexam = ReExamReminder.query.filter_by(
        student_id=current_user.id, is_active=True).all()
    return render_template('student_dashboard.html',
        reval_apps    = RevaluationApplication.query.filter_by(student_id=current_user.id)
                        .order_by(RevaluationApplication.applied_at.desc()).all(),
        photo_apps    = PhotocopyApplication.query.filter_by(student_id=current_user.id)
                        .order_by(PhotocopyApplication.applied_at.desc()).all(),
        notifs        = Notification.query.filter_by(user_id=current_user.id)
                        .order_by(Notification.created_at.desc()).limit(10).all(),
        unread        = Notification.query.filter_by(user_id=current_user.id, is_read=False).count(),
        result_uploads= ResultUpload.query.filter_by(student_id=current_user.id)
                        .order_by(ResultUpload.uploaded_at.desc()).limit(5).all(),
        active_rexam  = active_rexam,
        REVAL_FEE     = REVAL_FEE_PER_PAPER,
        PHOTO_FEE     = PHOTOCOPY_FEE
    )

# ─── Apply Revaluation ────────────────────────────────────────────────────────

@app.route('/student/apply/revaluation', methods=['GET','POST'])
@login_required
def apply_revaluation():
    if current_user.role == 'admin': return redirect(url_for('admin_dashboard'))
    if request.method == 'POST':
        num_papers = int(request.form.get('num_papers', 1))
        fee        = num_papers * REVAL_FEE_PER_PAPER

        marksheet_path = None
        ms_file = request.files.get('marksheet_file')
        if ms_file and ms_file.filename:
            if not allowed_file(ms_file.filename):
                flash('Marksheet must be PNG, JPG, or PDF.', 'error')
                return render_template('apply_revaluation.html', REVAL_FEE=REVAL_FEE_PER_PAPER, SEMESTER_SUBJECTS=SEMESTER_SUBJECTS)
            ms_name = f"ms_{current_user.id}_{datetime.utcnow().strftime('%Y%m%d%H%M%S')}_{secure_filename(ms_file.filename)}"
            ms_path = os.path.join(app.config['UPLOAD_FOLDER'], 'marksheets', ms_name)
            ms_file.save(ms_path)
            marksheet_path = ms_name
        else:
            flash('Please attach your previous semester marksheet.', 'error')
            return render_template('apply_revaluation.html', REVAL_FEE=REVAL_FEE_PER_PAPER, SEMESTER_SUBJECTS=SEMESTER_SUBJECTS)

        obj = RevaluationApplication(
            student_id    = current_user.id,
            subject_name  = request.form.get('subject_name') or request.form.get('subject_name_other','').strip(),
            exam_year     = request.form.get('exam_year'),
            semester      = request.form.get('semester'),
            prn_number    = request.form.get('prn_number'),
            num_papers    = num_papers,
            fee_amount    = fee,
            marksheet_file= marksheet_path,
            status        = 'Applied'
        )
        db.session.add(obj); db.session.commit()

        from email_utils import send_application_submitted
        from whatsapp_utils import wa_application_submitted
        notify_student(current_user,
            'Revaluation Applied',
            f'Revaluation for {obj.subject_name} submitted. Fee: ₹{fee}', 'info',
            email_fn=send_application_submitted,
            email_args=(current_user.name, current_user.email, 'Revaluation',
                        obj.subject_name, obj.id, fee),
            wa_fn=wa_application_submitted,
            wa_args=(current_user.name, 'Revaluation', obj.subject_name, obj.id, fee))

        flash(f'Revaluation applied! Fee payable: ₹{fee}. Email & WhatsApp sent.', 'success')
        return redirect(url_for('student_dashboard'))

    return render_template('apply_revaluation.html', REVAL_FEE=REVAL_FEE_PER_PAPER, SEMESTER_SUBJECTS=SEMESTER_SUBJECTS)

# ─── Apply Photocopy ──────────────────────────────────────────────────────────

@app.route('/student/apply/photocopy', methods=['GET','POST'])
@login_required
def apply_photocopy():
    if current_user.role == 'admin': return redirect(url_for('admin_dashboard'))
    if request.method == 'POST':
        num_papers = int(request.form.get('num_papers', 1))
        fee        = num_papers * PHOTOCOPY_FEE

        obj = PhotocopyApplication(
            student_id   = current_user.id,
            subject_name = request.form.get('subject_name') or request.form.get('subject_name_other','').strip(),
            exam_year    = request.form.get('exam_year'),
            semester     = request.form.get('semester'),
            prn_number   = request.form.get('prn_number'),
            num_papers   = num_papers,
            fee_amount   = fee,
            status       = 'Applied',
            expected_by  = datetime.utcnow() + timedelta(days=20)
        )
        db.session.add(obj); db.session.commit()

        from email_utils import send_application_submitted
        from whatsapp_utils import wa_application_submitted
        notify_student(current_user,
            'Photocopy Applied',
            f'Photocopy for {obj.subject_name} submitted. Fee: ₹{fee}', 'info',
            email_fn=send_application_submitted,
            email_args=(current_user.name, current_user.email, 'Photocopy',
                        obj.subject_name, obj.id, fee),
            wa_fn=wa_application_submitted,
            wa_args=(current_user.name, 'Photocopy', obj.subject_name, obj.id, fee))

        flash(f'Photocopy applied! Fee payable: ₹{fee}. Email & WhatsApp sent.', 'success')
        return redirect(url_for('student_dashboard'))

    return render_template('apply_photocopy.html', PHOTO_FEE=PHOTOCOPY_FEE, SEMESTER_SUBJECTS=SEMESTER_SUBJECTS)

# ─── Re-Exam ──────────────────────────────────────────────────────────────────

@app.route('/student/rexam/<int:reminder_id>/filled', methods=['POST'])
@login_required
def rexam_form_filled(reminder_id):
    rem = ReExamReminder.query.get_or_404(reminder_id)
    if rem.student_id != current_user.id:
        flash('Unauthorized.', 'error'); return redirect(url_for('student_dashboard'))
    rem.is_active      = False
    rem.form_filled_at = datetime.utcnow()
    db.session.commit()
    flash(f'Re-exam form for {rem.subject_name} marked as filled. Good luck! ', 'success')
    return redirect(url_for('student_dashboard'))

@app.route('/student/photocopy/<int:app_id>/mark_received', methods=['POST'])
@login_required
def mark_photocopy_received(app_id):
    obj = PhotocopyApplication.query.get_or_404(app_id)
    if obj.student_id != current_user.id:
        flash('Unauthorized.', 'error'); return redirect(url_for('student_dashboard'))
    obj.status = 'Delivered'; obj.received_at = datetime.utcnow()
    db.session.commit(); flash('Marked as received!', 'success')
    return redirect(url_for('student_dashboard'))

@app.route('/student/result/upload', methods=['GET','POST'])
@login_required
def upload_result():
    if current_user.role == 'admin': return redirect(url_for('admin_dashboard'))
    parsed = None; error = None
    if request.method == 'POST':
        file = request.files.get('result_file')
        if not file or file.filename == '':
            error = 'Please select a file.'
        elif not allowed_file(file.filename):
            error = 'Only PNG, JPG, JPEG, PDF files are allowed.'
        else:
            filename = secure_filename(file.filename)
            saved_path = os.path.join(app.config['UPLOAD_FOLDER'], filename)
            file.save(saved_path)
            try:
                from ocr_utils import process_upload
                parsed = process_upload(saved_path)
                obj = ResultUpload(student_id=current_user.id, filename=filename,
                    extracted_text=parsed.get('raw',''),
                    parsed_json=json.dumps(parsed),
                    sgpa=parsed.get('sgpa'), cgpa=parsed.get('cgpa'))
                db.session.add(obj); db.session.commit()
                flash('Result processed!', 'success')
            except Exception as e:
                error = str(e)
    return render_template('upload_result.html', parsed=parsed, error=error)

@app.route('/student/profile', methods=['GET','POST'])
@login_required
def student_profile():
    if current_user.role == 'admin': return redirect(url_for('admin_dashboard'))
    if request.method == 'POST':
        current_user.phone = request.form.get('phone','').strip()
        current_user.department = request.form.get('department','').strip()
        current_user.semester = request.form.get('semester','').strip()
        db.session.commit(); flash('Profile updated!', 'success')
    return render_template('student_profile.html')

@app.route('/student/notifications/read/<int:nid>', methods=['POST'])
@login_required
def mark_notification_read(nid):
    n = Notification.query.get_or_404(nid)
    if n.user_id == current_user.id:
        n.is_read = True; db.session.commit()
    return jsonify({'ok': True})

@app.route('/student/notifications/read_all', methods=['POST'])
@login_required
def mark_all_read():
    Notification.query.filter_by(user_id=current_user.id, is_read=False).update({'is_read': True})
    db.session.commit(); return jsonify({'ok': True})

# ─── Admin Dashboard ──────────────────────────────────────────────────────────

@app.route('/admin/dashboard')
@login_required
@admin_required
def admin_dashboard():
    failed_after_reval = ReExamReminder.query.filter_by(is_active=True).count()

    # ── Analytics data ────────────────────────────────────────────────────
    total_students = User.query.filter_by(role='student').count()
    all_reval = RevaluationApplication.query.all()

    # KT % (students who have at least one reval app)
    students_with_kt = db.session.query(RevaluationApplication.student_id)\
        .distinct().count()
    kt_percent = round((students_with_kt / total_students * 100) if total_students else 0, 1)

    # Subject-wise failure rate
    from sqlalchemy import func
    subject_counts = db.session.query(
        RevaluationApplication.subject_name,
        func.count(RevaluationApplication.id).label('total')
    ).group_by(RevaluationApplication.subject_name)\
     .order_by(func.count(RevaluationApplication.id).desc()).limit(8).all()

    subject_labels = [s.subject_name[:20] for s in subject_counts]
    subject_values = [s.total for s in subject_counts]

    # Semester-wise trend
    sem_counts = db.session.query(
        RevaluationApplication.semester,
        func.count(RevaluationApplication.id).label('total')
    ).group_by(RevaluationApplication.semester)\
     .order_by(RevaluationApplication.semester).all()

    sem_labels = [f"Sem {s.semester}" for s in sem_counts]
    sem_values = [s.total for s in sem_counts]

    # Department-wise (based on user department field)
    dept_counts = db.session.query(
        User.department,
        func.count(RevaluationApplication.id).label('total')
    ).join(RevaluationApplication, User.id == RevaluationApplication.student_id)\
     .filter(User.department != '')\
     .group_by(User.department).all()

    dept_labels = [d.department for d in dept_counts]
    dept_values = [d.total for d in dept_counts]

    # Weak students (most KT)
    weak_students = db.session.query(
        User.name, User.email, User.department, User.semester,
        func.count(RevaluationApplication.id).label('kt_count')
    ).join(RevaluationApplication, User.id == RevaluationApplication.student_id)\
     .group_by(User.id)\
     .order_by(func.count(RevaluationApplication.id).desc()).limit(10).all()

    return render_template('admin_dashboard.html',
        total_students     = total_students,
        pending_reval      = RevaluationApplication.query.filter_by(status='Applied').count(),
        pending_photo      = PhotocopyApplication.query.filter_by(status='Applied').count(),
        monitoring         = RevaluationApplication.query.filter_by(status='Monitoring').count(),
        failed_after_reval = failed_after_reval,
        recent_reval       = RevaluationApplication.query.order_by(
                                 RevaluationApplication.applied_at.desc()).limit(8).all(),
        recent_photo       = PhotocopyApplication.query.order_by(
                                 PhotocopyApplication.applied_at.desc()).limit(8).all(),
        # Analytics
        kt_percent         = kt_percent,
        students_with_kt   = students_with_kt,
        subject_labels     = json.dumps(subject_labels),
        subject_values     = json.dumps(subject_values),
        sem_labels         = json.dumps(sem_labels),
        sem_values         = json.dumps(sem_values),
        dept_labels        = json.dumps(dept_labels),
        dept_values        = json.dumps(dept_values),
        weak_students      = weak_students,
    )

@app.route('/admin/revaluation')
@login_required
@admin_required
def admin_reval_list():
    sf = request.args.get('status','')
    q  = RevaluationApplication.query
    if sf: q = q.filter_by(status=sf)
    return render_template('admin_reval_list.html',
        apps=q.order_by(RevaluationApplication.applied_at.desc()).all(),
        status_filter=sf)

@app.route('/admin/marksheet/<path:filename>')
@login_required
@admin_required
def view_marksheet(filename):
    return send_from_directory(
        os.path.join(app.config['UPLOAD_FOLDER'], 'marksheets'), filename)

@app.route('/uploads/caste_certs/<path:filename>')
@login_required
def view_caste_cert(filename):
    """Serve caste certificates to authenticated users"""
    return send_from_directory(
        os.path.join(app.config['UPLOAD_FOLDER'], 'caste_certs'), filename)

@app.route('/admin/revaluation/<int:app_id>/verify', methods=['POST'])
@login_required
@admin_required
def verify_reval(app_id):
    obj = RevaluationApplication.query.get_or_404(app_id)
    obj.status='Monitoring'; obj.verified_at=datetime.utcnow()
    obj.admin_notes=request.form.get('admin_notes','')
    db.session.commit()
    from email_utils import send_reval_verified
    from whatsapp_utils import wa_reval_verified
    notify_student(obj.student,
        'Revaluation Verified',
        f'Your revaluation for {obj.subject_name} is now being monitored.', 'success',
        email_fn=send_reval_verified,
        email_args=(obj.student.name, obj.student.email, obj.subject_name, obj.id, obj.fee_amount),
        wa_fn=wa_reval_verified,
        wa_args=(obj.student.name, obj.subject_name, obj.id, obj.fee_amount))
    flash('Verified! Student notified via Email & WhatsApp.', 'success')
    return redirect(url_for('admin_reval_list'))

@app.route('/admin/revaluation/<int:app_id>/update_result', methods=['POST'])
@login_required
@admin_required
def update_reval_result(app_id):
    obj = RevaluationApplication.query.get_or_404(app_id)
    updated_result = request.form.get('updated_result','')
    result_outcome = request.form.get('result_outcome','')

    obj.status            = 'Result Updated'
    obj.result_updated_at = datetime.utcnow()
    obj.updated_result    = updated_result
    db.session.commit()

    from email_utils import send_result_updated
    from whatsapp_utils import wa_result_updated
    notify_student(obj.student,
        'Revaluation Result Updated',
        f'Your result for {obj.subject_name}: {updated_result}', 'success',
        email_fn=send_result_updated,
        email_args=(obj.student.name, obj.student.email,
                    obj.subject_name, updated_result, obj.id),
        wa_fn=wa_result_updated,
        wa_args=(obj.student.name, obj.subject_name, updated_result, obj.id))

    if result_outcome == 'fail':
        existing = ReExamReminder.query.filter_by(
            student_id=obj.student_id, reval_app_id=obj.id, is_active=True).first()
        if not existing:
            rem = ReExamReminder(
                student_id   = obj.student_id,
                reval_app_id = obj.id,
                subject_name = obj.subject_name,
                semester     = obj.semester,
                exam_year    = obj.exam_year,
                prn_number   = obj.prn_number,
                is_active    = True
            )
            db.session.add(rem); db.session.commit()

            from email_utils import send_rexam_reminder
            from whatsapp_utils import wa_rexam_reminder
            notify_student(obj.student,
                ' Re-Exam Form Alert',
                f'You failed {obj.subject_name} after revaluation. Please fill the re-exam form immediately.', 'warning',
                email_fn=send_rexam_reminder,
                email_args=(obj.student.name, obj.student.email, obj.subject_name,
                            obj.semester, obj.prn_number, 1),
                wa_fn=wa_rexam_reminder,
                wa_args=(obj.student.name, obj.subject_name, obj.semester, obj.prn_number, 1))

        flash(f'Result updated as FAIL. Re-exam reminders activated for {obj.student.name}.', 'warning')
    else:
        flash('Result updated! Student notified via Email & WhatsApp.', 'success')

    return redirect(url_for('admin_reval_list'))

@app.route('/admin/photocopy')
@login_required
@admin_required
def admin_photo_list():
    sf = request.args.get('status','')
    q  = PhotocopyApplication.query
    if sf: q = q.filter_by(status=sf)
    return render_template('admin_photo_list.html',
        apps=q.order_by(PhotocopyApplication.applied_at.desc()).all(), status_filter=sf)

@app.route('/admin/photocopy/<int:app_id>/verify', methods=['POST'])
@login_required
@admin_required
def verify_photo(app_id):
    obj = PhotocopyApplication.query.get_or_404(app_id)
    obj.status='In Progress'; obj.verified_at=datetime.utcnow()
    obj.admin_notes=request.form.get('admin_notes','')
    db.session.commit()
    from email_utils import send_photocopy_verified
    from whatsapp_utils import wa_photocopy_verified
    expected_str = obj.expected_by.strftime('%d %b %Y') if obj.expected_by else 'TBD'
    notify_student(obj.student,
        'Photocopy Verified',
        f'Your photocopy for {obj.subject_name} is being processed.', 'success',
        email_fn=send_photocopy_verified,
        email_args=(obj.student.name, obj.student.email, obj.subject_name,
                    expected_str, obj.id, obj.fee_amount),
        wa_fn=wa_photocopy_verified,
        wa_args=(obj.student.name, obj.subject_name, expected_str, obj.id, obj.fee_amount))
    flash('Verified! Student notified via Email & WhatsApp.', 'success')
    return redirect(url_for('admin_photo_list'))

@app.route('/admin/students')
@login_required
@admin_required
def admin_students():
    return render_template('admin_students.html',
        students=User.query.filter_by(role='student').order_by(User.created_at.desc()).all())

@app.route('/admin/caste-certificates')
@login_required
@admin_required
def admin_caste_certificates():
    """Display all students who uploaded caste certificates"""
    students = User.query.filter(
        User.role == 'student',
        User.caste_certificate != ''
    ).order_by(User.created_at.desc()).all()
    
    return render_template('admin_caste_certificates.html', students=students)

@app.route('/admin/caste-certificates/<int:student_id>/download')
@login_required
@admin_required
def download_caste_certificate(student_id):
    """Download a student's caste certificate"""
    student = User.query.get_or_404(student_id)
    
    if not student.caste_certificate or student.role != 'student':
        flash('Certificate not found.', 'error')
        return redirect(url_for('admin_caste_certificates'))
    
    try:
        cert_path = os.path.join(app.config['UPLOAD_FOLDER'], 'caste_certs', 
                                 student.caste_certificate)
        if not os.path.exists(cert_path):
            flash('File not found on server.', 'error')
            return redirect(url_for('admin_caste_certificates'))
        
        return send_from_directory(
            os.path.join(app.config['UPLOAD_FOLDER'], 'caste_certs'),
            student.caste_certificate,
            as_attachment=True,
            download_name=f"caste_cert_{student.name.replace(' ', '_')}_{student.id}"
        )
    except Exception as e:
        flash(f'Error downloading certificate: {str(e)}', 'error')
        return redirect(url_for('admin_caste_certificates'))

@app.route('/admin/rexam')
@login_required
@admin_required
def admin_rexam_list():
    reminders = ReExamReminder.query.order_by(ReExamReminder.created_at.desc()).all()
    return render_template('admin_rexam_list.html', reminders=reminders)

@app.route('/admin/rexam/send-reminders', methods=['POST'])
@login_required
@admin_required
def send_rexam_reminders_api():
    """Send re-exam reminders immediately and return updated data"""
    from scheduler import send_rexam_reminders as send_reminders_task
    
    try:
        send_reminders_task(app)
        reminders = ReExamReminder.query.order_by(ReExamReminder.created_at.desc()).all()
        active_count = len([r for r in reminders if r.is_active])
        completed_count = len([r for r in reminders if not r.is_active])
        
        reminder_list = []
        for r in reminders:
            reminder_list.append({
                'id': r.id,
                'student_name': r.student.name if r.student else 'N/A',
                'student_email': r.student.email if r.student else 'N/A',
                'subject_name': r.subject_name,
                'semester': r.semester,
                'exam_year': r.exam_year,
                'prn_number': r.prn_number,
                'reminder_count': r.reminder_count,
                'last_reminded': r.last_reminded.strftime('%d %b %Y') if r.last_reminded else '—',
                'is_active': r.is_active,
                'form_filled_at': r.form_filled_at.strftime('%d %b %Y') if r.form_filled_at else '—'
            })
        
        return jsonify({
            'success': True,
            'message': 'Reminders sent successfully!',
            'active_count': active_count,
            'completed_count': completed_count,
            'reminders': reminder_list
        })
    except Exception as e:
        return jsonify({
            'success': False,
            'message': f'Error sending reminders: {str(e)}'
        }), 500

@app.route('/admin/trigger_check', methods=['POST'])
@login_required
@admin_required
def trigger_check():
    from scheduler import check_mu_results, photocopy_reminders, send_rexam_reminders
    check_mu_results(app); photocopy_reminders(app); send_rexam_reminders(app)
    flash('Manual check triggered (MU results + photocopy + re-exam reminders).', 'success')
    return redirect(url_for('admin_dashboard'))

# ─── Export CSV/Excel ─────────────────────────────────────────────────────────

@app.route('/admin/export/revaluation')
@login_required
@admin_required
def export_reval_csv():
    import csv, io
    apps = RevaluationApplication.query.order_by(RevaluationApplication.applied_at.desc()).all()
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(['ID','Student Name','Email','PRN','Subject','Semester','Exam Year',
                     'Papers','Fee (₹)','Status','Applied At','Verified At','Result'])
    for a in apps:
        writer.writerow([
            a.id, a.student.name, a.student.email, a.prn_number,
            a.subject_name, a.semester, a.exam_year,
            a.num_papers, a.fee_amount, a.status,
            a.applied_at.strftime('%d-%m-%Y %H:%M') if a.applied_at else '',
            a.verified_at.strftime('%d-%m-%Y %H:%M') if a.verified_at else '',
            a.updated_result or ''
        ])
    response = make_response(output.getvalue())
    response.headers['Content-Type'] = 'text/csv'
    response.headers['Content-Disposition'] = f'attachment; filename=revaluation_applications_{datetime.now().strftime("%Y%m%d")}.csv'
    return response

@app.route('/admin/export/photocopy')
@login_required
@admin_required
def export_photo_csv():
    import csv, io
    apps = PhotocopyApplication.query.order_by(PhotocopyApplication.applied_at.desc()).all()
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(['ID','Student Name','Email','PRN','Subject','Semester','Exam Year',
                     'Papers','Fee (₹)','Status','Applied At','Expected By','Received At'])
    for a in apps:
        writer.writerow([
            a.id, a.student.name, a.student.email, a.prn_number,
            a.subject_name, a.semester, a.exam_year,
            a.num_papers, a.fee_amount, a.status,
            a.applied_at.strftime('%d-%m-%Y %H:%M') if a.applied_at else '',
            a.expected_by.strftime('%d-%m-%Y') if a.expected_by else '',
            a.received_at.strftime('%d-%m-%Y') if a.received_at else ''
        ])
    response = make_response(output.getvalue())
    response.headers['Content-Type'] = 'text/csv'
    response.headers['Content-Disposition'] = f'attachment; filename=photocopy_applications_{datetime.now().strftime("%Y%m%d")}.csv'
    return response

# ─── Result PDF Analytics ─────────────────────────────────────────────────────

# Simple in-memory store for the last export (keyed by a short token)
_EXPORT_CACHE = {}

@app.route('/admin/result-analytics', methods=['GET', 'POST'])
@login_required
@admin_required
def result_analytics():
    results = None
    error   = None
    filter_subject = ''
    filter_sem     = ''

    if request.method == 'POST':
        pdf_file = request.files.get('result_pdf')
        filter_subject = request.form.get('filter_subject', '').strip()
        filter_sem     = request.form.get('filter_sem', '').strip()

        if not pdf_file or pdf_file.filename == '':
            error = 'Please select a result PDF file.'
            error = 'Please select a result PDF or image file.'
        elif not pdf_file.filename.lower().endswith('.pdf'):
            error = 'Please upload a PDF file. This parser works specifically with MU result register PDFs.'
        else:
            try:
                filename  = secure_filename(pdf_file.filename)
                save_path = os.path.join(app.config['UPLOAD_FOLDER'], 'analytics_' + filename)
                pdf_file.save(save_path)

                # ── Parse PDF using precise MU result register parser ─────────
                from ocr_utils import parse_mu_result_pdf
                parsed = parse_mu_result_pdf(save_path)

                exam_info    = parsed['exam_info']
                all_students = parsed['all_students']
                failed_stds  = parsed['failed']
                atkt_stds    = parsed.get('atkt', [])

                # Cross-reference reval apps for ATKT students too
                for s in atkt_stds:
                    s['reval_apps'] = []
                    if s.get('seat'):
                        s['reval_apps'] = RevaluationApplication.query.filter(
                            RevaluationApplication.prn_number == s['seat']
                        ).all()
                    if not s['reval_apps']:
                        for fs in s.get('failed_subjects', []):
                            apps = RevaluationApplication.query.filter(
                                RevaluationApplication.subject_name.ilike(
                                    f"%{fs['subject'][:20]}%")
                            ).all()
                            s['reval_apps'].extend(apps)
                        # deduplicate
                        seen = set()
                        deduped = []
                        for a in s['reval_apps']:
                            if a.id not in seen:
                                seen.add(a.id); deduped.append(a)
                        s['reval_apps'] = deduped

                # ── Optional subject/sem filter ───────────────────────────────
                if filter_subject:
                    filtered_failed = []
                    for s in failed_stds:
                        s_copy = dict(s)
                        s_copy['failed_subjects'] = [
                            f for f in s['failed_subjects']
                            if filter_subject.lower() in f['subject'].lower()
                        ]
                        if s_copy['failed_subjects']:
                            filtered_failed.append(s_copy)
                    failed_stds = filtered_failed

                # ── Build per-subject breakdown ───────────────────────────────
                # key = subject name, value = {failed: [rows], reval_count, filing_rate}
                subjects_affected = {}
                all_failed_rows   = []   # flat: one row per (student, subject)

                for student in failed_stds:
                    for fs in student['failed_subjects']:
                        sname = fs['subject']

                        # Match reval applications in DB by seat number (PRN) + subject
                        reval_apps = []
                        if student['seat']:
                            reval_apps = RevaluationApplication.query.filter(
                                RevaluationApplication.prn_number == student['seat']
                            ).all()
                            if filter_subject:
                                reval_apps = [a for a in reval_apps
                                              if filter_subject.lower() in a.subject_name.lower()]

                        # Fallback: match by subject name only
                        if not reval_apps:
                            reval_apps = RevaluationApplication.query.filter(
                                RevaluationApplication.subject_name.ilike(f'%{sname[:20]}%')
                            ).all()

                        if filter_sem:
                            reval_apps = [a for a in reval_apps
                                          if str(a.semester) == filter_sem]

                        row = {
                            'name':       student['name'],
                            'seat':       student['seat'],
                            'gender':     student['gender'],
                            'marks':      fs['marks'],
                            'grade':      fs['grade'],
                            'subject':    sname,
                            'total_marks': student['total_marks'],
                            'all_subjects': student.get('all_subjects', []),
                            'reval_apps': reval_apps,
                        }

                        if sname not in subjects_affected:
                            subjects_affected[sname] = {
                                'failed': [], 'reval_count': 0, 'filing_rate': 0
                            }
                        subjects_affected[sname]['failed'].append(row)
                        if reval_apps:
                            subjects_affected[sname]['reval_count'] += 1
                        all_failed_rows.append(row)

                # Filing rates per subject
                for sname, data in subjects_affected.items():
                    total = len(data['failed'])
                    data['filing_rate'] = round(data['reval_count'] / total * 100) if total else 0

                total_failed  = len(failed_stds)          # unique failed students
                total_fail_entries = len(all_failed_rows)  # student×subject rows
                reval_filed   = sum(1 for r in all_failed_rows if r['reval_apps'])
                reval_pending = total_fail_entries - reval_filed

                subject_fail_counts = [(s, len(d['failed'])) for s, d in subjects_affected.items()]
                subject_fail_counts.sort(key=lambda x: x[1], reverse=True)

                # Store for export and save-to-db
                token = secrets.token_urlsafe(8)
                _EXPORT_CACHE[token] = {
                    'rows':    all_failed_rows,
                    'students': failed_stds,
                    'filename': filename,
                }
                if len(_EXPORT_CACHE) > 20:
                    del _EXPORT_CACHE[list(_EXPORT_CACHE.keys())[0]]

                results = {
                    'exam_info':           exam_info,
                    'total_students':      len(all_students),
                    'total_failed':        total_failed,
                    'total_fail_entries':  total_fail_entries,
                    'atkt_count':          len(atkt_stds),
                    'atkt_students':       atkt_stds,
                    'reval_filed':         reval_filed,
                    'reval_pending':       reval_pending,
                    'subjects_affected':   subjects_affected,
                    'all_failed':          all_failed_rows,
                    'failed_students':     failed_stds,
                    'subject_fail_counts': subject_fail_counts,
                    'export_token':        token,
                }

            except Exception as e:
                import traceback
                app.logger.error(f"[Result Analytics] {traceback.format_exc()}")
                error = f'Failed to process the PDF: {e}'

    return render_template('admin_result_analytics.html',
        results=results, error=error,
        filter_subject=filter_subject, filter_sem=filter_sem)


@app.route('/admin/result-analytics/export')
@login_required
@admin_required
def result_analytics_export():
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, GradientFill
    from openpyxl.utils import get_column_letter

    token    = request.args.get('data', '')
    semester = request.args.get('sem', '')   # admin-entered semester
    cache    = _EXPORT_CACHE.get(token, {})
    rows     = cache if isinstance(cache, list) else cache.get('rows', [])

    # Tag every row with the admin-confirmed semester for grouping
    if semester:
        for r in rows:
            r['semester'] = semester

    wb = Workbook()
    ws = wb.active
    ws.title = 'Failed Students'

    # ── Styles ────────────────────────────────────────────────────────────────
    header_font   = Font(name='Arial', bold=True, color='FFFFFF', size=11)
    header_fill   = PatternFill('solid', start_color='1A8A7A')   # teal
    subhdr_font   = Font(name='Arial', bold=True, color='FFFFFF', size=10)
    subhdr_fill   = PatternFill('solid', start_color='136B5E')   # dark teal
    fail_fill     = PatternFill('solid', start_color='FCE8E8')   # light red
    filed_fill    = PatternFill('solid', start_color='E8F5F3')   # light teal
    alt_fill      = PatternFill('solid', start_color='F5FAFA')   # very light teal
    thin          = Side(style='thin', color='C8D8D5')
    border        = Border(left=thin, right=thin, top=thin, bottom=thin)
    center        = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left          = Alignment(horizontal='left',   vertical='center', wrap_text=True)
    bold_red      = Font(name='Arial', bold=True, color='E05252', size=10)
    bold_teal     = Font(name='Arial', bold=True, color='136B5E', size=10)
    normal        = Font(name='Arial', size=10)
    bold_normal   = Font(name='Arial', bold=True, size=10)

    def hdr_cell(cell, value):
        cell.value     = value
        cell.font      = header_font
        cell.fill      = header_fill
        cell.border    = border
        cell.alignment = center

    def style_cell(cell, value, fnt=None, fill=None, align=None):
        cell.value     = value
        cell.font      = fnt   or normal
        cell.fill      = fill  or PatternFill()
        cell.border    = border
        cell.alignment = align or left

    # ── Title row ─────────────────────────────────────────────────────────────
    ws.merge_cells('A1:H1')
    title_cell = ws['A1']
    title_cell.value     = f'APSIT — Failed Students Report  ({datetime.now().strftime("%d %b %Y")})'
    title_cell.font      = Font(name='Arial', bold=True, size=13, color='136B5E')
    title_cell.fill      = PatternFill('solid', start_color='E8F5F3')
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 26

    # ── Column headers ────────────────────────────────────────────────────────
    headers = ['#', 'Student Name', 'Seat / PRN', 'Semester',
               'Subject (Failed)', 'Marks', 'Grade', 'Reval Filed?']
    for col, h in enumerate(headers, 1):
        hdr_cell(ws.cell(row=2, column=col), h)
    ws.row_dimensions[2].height = 22

    # ── Group rows by semester ─────────────────────────────────────────────────
    from itertools import groupby

    # Sort by semester first
    def sem_key(r):
        try:    return int(r.get('semester', '') or 0)
        except: return 0

    sorted_rows = sorted(rows, key=sem_key)

    data_row   = 3
    global_idx = 1

    for sem, group in groupby(sorted_rows, key=lambda r: r.get('semester', '')):
        group = list(group)

        # Semester group header
        ws.merge_cells(f'A{data_row}:H{data_row}')
        sem_label = f'  Semester {sem}' if sem else '  Semester — (not specified)'
        sem_cell  = ws[f'A{data_row}']
        sem_cell.value     = f'{sem_label}    ({len(group)} record(s))'
        sem_cell.font      = subhdr_font
        sem_cell.fill      = subhdr_fill
        sem_cell.alignment = Alignment(horizontal='left', vertical='center')
        ws.row_dimensions[data_row].height = 20
        data_row += 1

        for i, r in enumerate(group):
            fill = alt_fill if i % 2 == 0 else PatternFill()
            filed = bool(r.get('reval_apps'))

            style_cell(ws.cell(data_row, 1), global_idx,         normal,      fill, center)
            style_cell(ws.cell(data_row, 2), r.get('name',''),   bold_normal, fill, left)
            style_cell(ws.cell(data_row, 3), r.get('seat',''),   normal,      fill, center)
            style_cell(ws.cell(data_row, 4), sem or '—',         normal,      fill, center)
            style_cell(ws.cell(data_row, 5), r.get('subject',''),normal,      fail_fill if not filed else fill, left)
            style_cell(ws.cell(data_row, 6), r.get('marks',''),  bold_red,    fill, center)
            style_cell(ws.cell(data_row, 7), r.get('grade',''),  bold_red,    fill, center)

            if filed:
                reval_txt = f" Yes  (App #{', #'.join(str(a.id) for a in r['reval_apps'])})"
                style_cell(ws.cell(data_row, 8), reval_txt, bold_teal, filed_fill, left)
            else:
                style_cell(ws.cell(data_row, 8), ' Not Filed', Font(name='Arial', bold=True, color='E05252', size=10), fail_fill, center)

            ws.row_dimensions[data_row].height = 18
            data_row  += 1
            global_idx += 1

        data_row += 1   # blank spacer between semesters

    # ── Summary sheet ─────────────────────────────────────────────────────────
    ws2 = wb.create_sheet('Summary')
    ws2.column_dimensions['A'].width = 28
    ws2.column_dimensions['B'].width = 14

    summary_data = [
        ('Total Failed Records',     len(rows)),
        ('Revaluation Filed',        sum(1 for r in rows if r.get('reval_apps'))),
        ('Revaluation NOT Filed',    sum(1 for r in rows if not r.get('reval_apps'))),
        ('Unique Subjects',          len(set(r.get('subject','') for r in rows))),
        ('Generated On',             datetime.now().strftime('%d %b %Y %H:%M')),
    ]

    ws2['A1'] = 'Summary'
    ws2['A1'].font = Font(name='Arial', bold=True, size=12, color='136B5E')
    ws2['A1'].fill = PatternFill('solid', start_color='E8F5F3')
    ws2.merge_cells('A1:B1')
    ws2['A1'].alignment = Alignment(horizontal='center')

    for r_idx, (label, val) in enumerate(summary_data, 2):
        ws2.cell(r_idx, 1, label).font  = Font(name='Arial', bold=True, size=10)
        ws2.cell(r_idx, 2, val).font    = Font(name='Arial', size=10)
        ws2.cell(r_idx, 1).border       = border
        ws2.cell(r_idx, 2).border       = border

    # ── Column widths (main sheet) ─────────────────────────────────────────────
    col_widths = [5, 28, 14, 10, 36, 8, 8, 32]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # Freeze panes below title + header
    ws.freeze_panes = 'A3'

    # ── Stream response ────────────────────────────────────────────────────────
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    fname    = f'APSIT_Failed_Students_{datetime.now().strftime("%Y%m%d_%H%M")}.xlsx'
    response = make_response(buf.read())
    response.headers['Content-Type']        = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    response.headers['Content-Disposition'] = f'attachment; filename={fname}'
    return response


# ─── Failed Student Records — Save from PDF ───────────────────────────────────

@app.route('/admin/result-analytics/save', methods=['POST'])
@login_required
@admin_required
def save_failed_records():
    """
    Save/upsert failed student rows (from the in-memory export cache) to DB.
    Payload: { token: str, semester: str, exam_year: str }
    """
    from models import FailedStudentRecord
    data       = request.get_json(force=True)
    token      = data.get('token', '')
    semester   = (data.get('semester') or '').strip()
    exam_year  = (data.get('exam_year') or '').strip()
    pdf_source = (data.get('pdf_source') or '').strip()

    rows = _EXPORT_CACHE.get(token, {})
    if isinstance(rows, dict):
        rows = rows.get('rows', [])
    if not rows:
        return jsonify({'ok': False, 'msg': 'Session expired — please re-upload the PDF.'}), 400

    if not semester:
        return jsonify({'ok': False, 'msg': 'Please enter a semester number before saving.'}), 400

    saved = updated = 0
    for r in rows:
        # Use admin-entered semester strictly — PDF rows rarely carry semester
        rec_semester = semester
        rec_year     = exam_year

        matched_user = None
        if r.get('seat'):
            matched_user = User.query.filter(
                User.name.ilike(f"%{r['seat']}%")
            ).first()
        if not matched_user and r.get('name'):
            name_parts = r['name'].strip().split()
            if len(name_parts) >= 2:
                matched_user = User.query.filter(
                    User.name.ilike(f"%{name_parts[0]}%")
                ).filter(
                    User.name.ilike(f"%{name_parts[-1]}%")
                ).first()

        email_val = matched_user.email if matched_user else ''

        # Upsert key: student + subject + admin-entered semester + year
        existing = FailedStudentRecord.query.filter_by(
            student_name = r['name'],
            subject_name = r['subject'],
            semester     = rec_semester,
            exam_year    = rec_year,
        ).first()

        if existing:
            existing.marks       = r.get('marks', existing.marks)
            existing.grade       = r.get('grade', existing.grade)
            existing.seat_number = r.get('seat', existing.seat_number)
            existing.email       = email_val or existing.email
            existing.pdf_source  = pdf_source or existing.pdf_source
            existing.user_id     = matched_user.id if matched_user else existing.user_id
            updated += 1
        else:
            rec = FailedStudentRecord(
                student_name = r['name'],
                seat_number  = r.get('seat', ''),
                email        = email_val,
                semester     = rec_semester,
                exam_year    = rec_year,
                subject_name = r['subject'],
                marks        = r.get('marks', ''),
                grade        = r.get('grade', ''),
                pdf_source   = pdf_source,
                user_id      = matched_user.id if matched_user else None,
            )
            db.session.add(rec)
            saved += 1

    db.session.commit()
    return jsonify({'ok': True, 'saved': saved, 'updated': updated,
                    'total': saved + updated})


# ─── Failed Student Records — Browse / Manage ─────────────────────────────────

@app.route('/admin/failed-students')
@login_required
@admin_required
def failed_students_list():
    """Browse saved failed student records with filters."""
    from models import FailedStudentRecord
    sem_filter   = request.args.get('semester', '').strip()
    subj_filter  = request.args.get('subject', '').strip()
    year_filter  = request.args.get('exam_year', '').strip()
    filed_filter = request.args.get('filed', '').strip()

    q = FailedStudentRecord.query
    if sem_filter:
        q = q.filter_by(semester=sem_filter)
    if subj_filter:
        q = q.filter(FailedStudentRecord.subject_name.ilike(f'%{subj_filter}%'))
    if year_filter:
        q = q.filter_by(exam_year=year_filter)
    if filed_filter == 'yes':
        q = q.filter_by(exam_form_filed=True)
    elif filed_filter == 'no':
        q = q.filter_by(exam_form_filed=False)

    records = q.order_by(
        FailedStudentRecord.semester,
        FailedStudentRecord.student_name,
        FailedStudentRecord.subject_name,
    ).all()

    total     = len(records)
    filed     = sum(1 for r in records if r.exam_form_filed)
    not_filed = total - filed
    semesters = sorted(set(r.semester for r in records if r.semester))
    subjects  = sorted(set(r.subject_name for r in records))
    exam_years = sorted(set(r.exam_year for r in records if r.exam_year))

    return render_template('admin_failed_students.html',
        records=records, total=total, filed=filed, not_filed=not_filed,
        semesters=semesters, subjects=subjects, exam_years=exam_years,
        sem_filter=sem_filter, subj_filter=subj_filter,
        year_filter=year_filter, filed_filter=filed_filter)


@app.route('/admin/failed-students/export-excel')
@login_required
@admin_required
def export_failed_students_excel():
    """Export saved FailedStudentRecord rows from DB to a formatted Excel file."""
    import io
    from itertools import groupby
    from models import FailedStudentRecord
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    # ── Apply same filters as the list view ───────────────────────────────────
    sem_filter   = request.args.get('semester', '').strip()
    subj_filter  = request.args.get('subject', '').strip()
    year_filter  = request.args.get('exam_year', '').strip()
    filed_filter = request.args.get('filed', '').strip()

    q = FailedStudentRecord.query
    if sem_filter:   q = q.filter_by(semester=sem_filter)
    if subj_filter:  q = q.filter(FailedStudentRecord.subject_name.ilike(f'%{subj_filter}%'))
    if year_filter:  q = q.filter_by(exam_year=year_filter)
    if filed_filter == 'yes': q = q.filter_by(exam_form_filed=True)
    elif filed_filter == 'no': q = q.filter_by(exam_form_filed=False)

    records = q.order_by(
        FailedStudentRecord.semester,
        FailedStudentRecord.student_name,
        FailedStudentRecord.subject_name,
    ).all()

    # ── Build workbook ────────────────────────────────────────────────────────
    wb = Workbook()
    ws = wb.active
    ws.title = 'Failed Students'

    thin        = Side(style='thin', color='C8D8D5')
    border      = Border(left=thin, right=thin, top=thin, bottom=thin)
    center      = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left_align  = Alignment(horizontal='left',   vertical='center', wrap_text=True)

    hdr_font    = Font(name='Arial', bold=True, color='FFFFFF', size=11)
    hdr_fill    = PatternFill('solid', start_color='1A8A7A')
    sem_font    = Font(name='Arial', bold=True, color='FFFFFF', size=10)
    sem_fill    = PatternFill('solid', start_color='136B5E')
    alt_fill    = PatternFill('solid', start_color='F5FAFA')
    fail_fill   = PatternFill('solid', start_color='FCE8E8')
    filed_fill  = PatternFill('solid', start_color='E8F5F3')
    normal      = Font(name='Arial', size=10)
    bold_n      = Font(name='Arial', bold=True, size=10)
    red_bold    = Font(name='Arial', bold=True, color='E05252', size=10)
    teal_bold   = Font(name='Arial', bold=True, color='136B5E', size=10)

    # Title
    ws.merge_cells('A1:J1')
    tc = ws['A1']
    tc.value     = f'APSIT — Failed Student Records  ({datetime.now().strftime("%d %b %Y")})'
    tc.font      = Font(name='Arial', bold=True, size=13, color='136B5E')
    tc.fill      = PatternFill('solid', start_color='E8F5F3')
    tc.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 26

    # Header row
    headers = ['#', 'Student Name', 'Seat / PRN', 'Semester', 'Exam Year',
               'Subject', 'Marks', 'Grade', 'Exam Form Filed?', 'Email on Record']
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=2, column=col)
        c.value     = h
        c.font      = hdr_font
        c.fill      = hdr_fill
        c.border    = border
        c.alignment = center
    ws.row_dimensions[2].height = 22

    data_row = 3
    idx      = 1

    for sem, group in groupby(records, key=lambda r: r.semester):
        group = list(group)

        # Semester separator
        ws.merge_cells(f'A{data_row}:J{data_row}')
        sc = ws[f'A{data_row}']
        sc.value     = f'  Semester {sem or "—"}    ({len(group)} record(s))'
        sc.font      = sem_font
        sc.fill      = sem_fill
        sc.alignment = Alignment(horizontal='left', vertical='center')
        ws.row_dimensions[data_row].height = 20
        data_row += 1

        for i, rec in enumerate(group):
            fill        = alt_fill if i % 2 == 0 else PatternFill()
            is_filed    = rec.exam_form_filed
            email_val   = rec.email or (rec.user.email if rec.user else '')
            s_fill      = filed_fill if is_filed else fail_fill

            def sc(col, val, fnt=None, fl=None, al=None):
                c = ws.cell(data_row, col)
                c.value     = val
                c.font      = fnt  or normal
                c.fill      = fl   or fill
                c.border    = border
                c.alignment = al   or left_align

            sc(1,  idx,                         normal,    fill,    center)
            sc(2,  rec.student_name,             bold_n,    fill,    left_align)
            sc(3,  rec.seat_number or '—',       normal,    fill,    center)
            sc(4,  f'Sem {rec.semester}' if rec.semester else '—', normal, fill, center)
            sc(5,  rec.exam_year or '—',         normal,    fill,    center)
            sc(6,  rec.subject_name,             normal,    s_fill,  left_align)
            sc(7,  f'{rec.marks} / {rec.grade}' if rec.grade else rec.marks, red_bold, fill, center)
            sc(8,  rec.grade or '—',             red_bold,  fill,    center)
            sc(9,  ' Filed' if is_filed else ' Not Filed',
                   teal_bold if is_filed else red_bold, filed_fill if is_filed else fail_fill, center)
            sc(10, email_val or '—',             normal,    fill,    left_align)

            ws.row_dimensions[data_row].height = 18
            data_row += 1
            idx      += 1

        data_row += 1  # blank row between semesters

    # Column widths
    for i, w in enumerate([5, 28, 14, 10, 14, 34, 12, 8, 14, 28], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A3'

    # Summary sheet
    ws2 = wb.create_sheet('Summary')
    ws2['A1'] = 'Summary'
    ws2['A1'].font = Font(name='Arial', bold=True, size=12, color='136B5E')
    ws2['A1'].fill = PatternFill('solid', start_color='E8F5F3')
    ws2.merge_cells('A1:B1')
    ws2['A1'].alignment = Alignment(horizontal='center')
    ws2.column_dimensions['A'].width = 28
    ws2.column_dimensions['B'].width = 16

    summary = [
        ('Total Records',           len(records)),
        ('Exam Form Filed',         sum(1 for r in records if r.exam_form_filed)),
        ('Form NOT Filed',          sum(1 for r in records if not r.exam_form_filed)),
        ('Semester Filter',         sem_filter  or 'All'),
        ('Exam Year Filter',        year_filter or 'All'),
        ('Subject Filter',          subj_filter or 'All'),
        ('Generated On',            datetime.now().strftime('%d %b %Y %H:%M')),
    ]
    for r_i, (label, val) in enumerate(summary, 2):
        ws2.cell(r_i, 1, label).font  = Font(name='Arial', bold=True, size=10)
        ws2.cell(r_i, 2, val).font    = Font(name='Arial', size=10)
        ws2.cell(r_i, 1).border       = border
        ws2.cell(r_i, 2).border       = border

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    fname    = f'APSIT_FailedRecords_DB_{datetime.now().strftime("%Y%m%d_%H%M")}.xlsx'
    response = make_response(buf.read())
    response.headers['Content-Type']        = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    response.headers['Content-Disposition'] = f'attachment; filename={fname}'
    return response


@app.route('/admin/failed-students/<int:rec_id>/toggle-filed', methods=['POST'])
@login_required
@admin_required
def toggle_form_filed(rec_id):
    from models import FailedStudentRecord
    rec = FailedStudentRecord.query.get_or_404(rec_id)
    rec.exam_form_filed = not rec.exam_form_filed
    rec.form_filed_at   = datetime.utcnow() if rec.exam_form_filed else None
    db.session.commit()
    return jsonify({'ok': True, 'filed': rec.exam_form_filed})


# ─── Send Reminder Emails to Non-Filers ───────────────────────────────────────

@app.route('/admin/failed-students/send-reminders', methods=['POST'])
@login_required
@admin_required
def send_exam_reminders():
    """
    Send reminder emails to all (or filtered) students who have NOT filed
    the exam form yet and who have an email address on record.
    """
    from models import FailedStudentRecord
    from email_utils import send_exam_form_reminder

    data        = request.get_json(force=True)
    sem_filter  = (data.get('semester') or '').strip()
    subj_filter = (data.get('subject') or '').strip()
    rec_ids     = data.get('ids', [])  # optional: specific IDs

    q = FailedStudentRecord.query.filter_by(exam_form_filed=False)
    if sem_filter:
        q = q.filter_by(semester=sem_filter)
    if subj_filter:
        q = q.filter(FailedStudentRecord.subject_name.ilike(f'%{subj_filter}%'))
    if rec_ids:
        q = q.filter(FailedStudentRecord.id.in_(rec_ids))

    targets = q.all()

    sent = skipped_no_email = 0
    for rec in targets:
        email_addr = rec.email or (rec.user.email if rec.user else '')
        if not email_addr:
            skipped_no_email += 1
            continue

        name = rec.student_name
        if rec.user:
            name = rec.user.name  # prefer DB name

        rec.mail_sent_count += 1
        rec.last_mail_sent   = datetime.utcnow()
        db.session.flush()

        send_exam_form_reminder(
            mail,
            name         = name,
            email        = email_addr,
            subject      = rec.subject_name,
            semester     = rec.semester,
            exam_year    = rec.exam_year,
            seat         = rec.seat_number,
            mail_count   = rec.mail_sent_count,
        )
        sent += 1

    db.session.commit()
    return jsonify({
        'ok': True,
        'sent': sent,
        'skipped_no_email': skipped_no_email,
        'total_targets': len(targets),
    })


# ─── Init DB ──────────────────────────────────────────────────────────────────

def init_db():
    with app.app_context():
        db.create_all()
        if not User.query.filter_by(email='admin@mu.ac.in').first():
            admin = User(name='Admin', email='admin@mu.ac.in', role='admin')
            admin.set_password('admin123')
            db.session.add(admin); db.session.commit()
            print(" Default admin: admin@mu.ac.in / admin123")

if __name__ == '__main__':
    init_db()
    from scheduler import start_scheduler
    start_scheduler(app)
    app.run(debug=True, use_reloader=False)